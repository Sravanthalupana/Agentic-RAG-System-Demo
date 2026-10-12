from io import BytesIO

import pytest
import tiktoken
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from rag.documents import MAX_FILE_BYTES, chunk_sections, load_document
from rag.models import Section


def pdf_bytes(text: str) -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
    )
    stream = DecodedStreamObject()
    stream.set_data(f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode())
    page[NameObject("/Contents")] = stream
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("The support window is 9 AM to 5 PM.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Escalation"
    table.cell(0, 1).text = "Contact Mira"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_load_formats_and_locations() -> None:
    assert load_document("../notes.TXT", b"\xef\xbb\xbfHello")[0] == Section(
        "notes.TXT", "text", "Hello"
    )
    pdf = load_document("guide.pdf", pdf_bytes("Warranty lasts 2 years."))
    assert pdf[0].location == "page 1"
    assert "2 years" in pdf[0].text
    docx = load_document("guide.docx", docx_bytes())
    assert docx[0].location == "block 1"
    assert "Contact Mira" in docx[1].text


@pytest.mark.parametrize(
    "filename,data",
    [
        ("x.exe", b"x"),
        ("x.txt", b""),
        ("x.txt", b"\xff"),
        ("x.pdf", b"invalid"),
        ("x.docx", b"invalid"),
        ("x.txt", b"   "),
    ],
)
def test_reject_invalid(filename: str, data: bytes) -> None:
    with pytest.raises(ValueError):
        load_document(filename, data)


def test_size_limit_and_scanned_pdf() -> None:
    with pytest.raises(ValueError, match="20 MB"):
        load_document("large.txt", b"x" * (MAX_FILE_BYTES + 1))
    writer = PdfWriter()
    writer.add_blank_page(width=10, height=10)
    output = BytesIO()
    writer.write(output)
    with pytest.raises(ValueError, match="OCR"):
        load_document("scan.pdf", output.getvalue())


def test_chunk_boundaries_overlap_and_ids() -> None:
    section = Section("a.txt", "text", "word " * 50)
    chunks = chunk_sections([section], 12, 3)
    encoding = tiktoken.get_encoding("cl100k_base")
    assert len(chunks) > 1
    assert all(len(encoding.encode(chunk.text)) <= 12 for chunk in chunks)
    assert encoding.encode(chunks[0].text)[-3:] == encoding.encode(chunks[1].text)[:3]
    assert chunks == chunk_sections([section], 12, 3)
    assert len({chunk.id for chunk in chunks}) == len(chunks)


def test_unicode_and_special_tokens() -> None:
    text = "世界🌍 " * 20 + "<|endoftext|>"
    chunks = chunk_sections([Section("a.txt", "text", text)], 12, 0)
    assert "".join(chunk.text for chunk in chunks) == text
    assert all("\ufffd" not in chunk.text for chunk in chunks)


def test_invalid_chunk_parameters() -> None:
    with pytest.raises(ValueError):
        chunk_sections([], 10, 10)
